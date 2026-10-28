"""卷面排版（範本 → 新卷）的後端測試。

獨立成一個檔案：只依賴 server.Application 的公開介面（dispatch），
不碰 test_server 的測試夾具。
"""
import base64
import io
import json
import tempfile
import unittest
from urllib.parse import urlencode
from unittest.mock import patch

from docx import Document
from docx.shared import Pt

import latex_export
import omml
import paper_format
import server
from docx.enum.text import WD_ALIGN_PARAGRAPH

DOCX_MIME = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'


def make_docx(blocks):
    """blocks: [(文字, 字級, 粗體)]，回傳 docx bytes。"""
    doc = Document()
    for text, size, bold in blocks:
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.font.size = Pt(size)
        run.font.bold = bold
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def data_url(data):
    return 'data:' + DOCX_MIME + ';base64,' + base64.b64encode(data).decode()


TEMPLATE = [
    ('某某中學 數學測驗', 16, True),
    ('一、選擇題：(每題 4 分)', 12, True),
    ('1.　下列哪一個是質數？（　　）', 12, False),
    ('A.　4\tB.　6', 12, False),
]


def content_docx():
    doc = Document()
    doc.add_paragraph('中三數學 小測')
    doc.add_paragraph('(本題目卷共兩頁，禁止使用計算機)')
    p = doc.add_paragraph('一、計算題')
    q = doc.add_paragraph('1.　解方程 ')
    q._p.append(omml.latex_to_omml('x^{2}-5x+6=0'))
    q.add_run('。')
    q2 = doc.add_paragraph('2.　計算 ')
    q2._p.append(omml.latex_to_omml(r'\frac{1}{2}+\frac{1}{3}'))
    q2.add_run('。')
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# 真實卷常見的寫法：選項那一行夾著方程式，還用 tab 分欄。
# 這些是範本自己的「範例內容」，套版後必須一個字都不留。
LEFTOVER_FORMULAS = ['x^{2}-2x+1=0', r'x-2y=\frac{1}{3}', '3xy+6=5xy', r'\frac{1}{x}+4=x^{2}']


def make_docx_with_math_option():
    doc = Document()
    doc.add_paragraph('某某中學 數學測驗')
    doc.add_paragraph('一、選擇題：（每題 4 分）')
    q = doc.add_paragraph()
    q.add_run('1.　下列哪一個是質數？')
    q._p.append(omml.latex_to_omml('x^{2}-5x+6=0'))
    q.add_run('（　　）')
    p = doc.add_paragraph()
    p.add_run('A. ')
    for tex in LEFTOVER_FORMULAS:
        p._p.append(omml.latex_to_omml(tex))
        p.add_run('\t')
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def content_with_options():
    doc = Document()
    doc.add_paragraph('中三數學 小測')
    doc.add_paragraph('一、選擇題')
    doc.add_paragraph('1.　下列哪一個是質數？')
    for option in ('A.　4', 'B.　6', 'C.　9', 'D.　11'):
        doc.add_paragraph(option)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_docx_with_left_subtitle():
    """範本把副標題靠左——套版後卷頭要整塊置中。"""
    doc = Document()
    title = doc.add_paragraph('某某中學 數學測驗')
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    sub = doc.add_paragraph('(禁止使用計算機)')
    sub.alignment = WD_ALIGN_PARAGRAPH.LEFT
    doc.add_paragraph('一、選擇題：(每題 4 分)')
    doc.add_paragraph('1.　下列哪一個是質數？（　　）')
    doc.add_paragraph('A.　4\tB.　6')
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


class PaperTemplateTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.app = server.Application(self.tmp.name, server.ROOT)
        self.app.initialize()
        with self.app.db() as db:
            teacher = db.execute("SELECT * FROM users WHERE role='teacher' ORDER BY id LIMIT 1").fetchone()
            student = db.execute("SELECT * FROM users WHERE role='student' ORDER BY id LIMIT 1").fetchone()
        self.teacher = dict(teacher) if teacher else None
        self.student = dict(student) if student else None
        if not self.teacher:
            self.skipTest('沒有教師帳號')

    def tearDown(self):
        self.tmp.cleanup()

    def call(self, method, path, body=None, user=None, expect=200):
        try:
            payload, _ = self.app.dispatch(method, path, body or {}, user or self.teacher,
                                           None, '127.0.0.1')
        except server.APIError as exc:
            self.assertEqual(exc.status, expect, exc.message)
            return {'error': exc.message}
        self.assertEqual(expect, 200)
        return payload

    def upload(self, name='培華範本'):
        return self.call('POST', '/api/paper-templates',
                         {'name': name, 'file': data_url(make_docx(TEMPLATE))})

    def test_upload_list_and_spec(self):
        created = self.upload()
        template = created['template']
        self.assertEqual(template['name'], '培華範本')
        self.assertEqual(template['spec']['marksPosition'], 'none')
        self.assertEqual(template['spec']['optionsPerLine'], 2)

        listed = self.call('GET', '/api/paper-templates')
        self.assertIn(template['id'], [t['id'] for t in listed['templates']])

    def test_apply_returns_docx_with_formulas(self):
        template = self.upload()['template']
        out = self.call('POST', f"/api/paper-templates/{template['id']}/apply",
                        {'file': data_url(content_docx())})
        self.assertEqual(out['stats']['questions'], 2)
        self.assertEqual(out['stats']['formulas'], 2)
        self.assertTrue(out['docx'].startswith('data:' + DOCX_MIME))
        raw = base64.b64decode(out['docx'].split(',', 1)[1])
        doc = Document(io.BytesIO(raw))
        equations = sum(1 for p in doc.paragraphs for c in p._p if c.tag.endswith('}oMath'))
        self.assertEqual(equations, 2, '公式要原樣帶進輸出')
        self.assertEqual(doc.paragraphs[0].text, '中三數學 小測', '卷名要換成新卷的')

    def test_template_example_formulas_do_not_leak_into_output(self):
        """範本的選項行夾著範例公式時，套版後不能拖在每一行後面。"""
        template = self.call('POST', '/api/paper-templates',
                             {'name': '有公式的範本',
                              'file': data_url(make_docx_with_math_option())})['template']
        out = self.call('POST', f"/api/paper-templates/{template['id']}/apply",
                        {'file': data_url(content_with_options())})
        raw = base64.b64decode(out['docx'].split(',', 1)[1])
        doc = Document(io.BytesIO(raw))

        for p in doc.paragraphs:
            self.assertNotIn('3xy', p._p.xml, f'範本的範例公式漏出來了：{p.text}')
            self.assertNotIn('x-2y', p._p.xml, f'範本的範例公式漏出來了：{p.text}')

        options = [p for p in doc.paragraphs
                   if p.text.strip()[:2] in ('A.', 'B.', 'C.', 'D.')]
        self.assertEqual(len(options), 1, '選項要打橫排在同一行')
        line = options[0].text
        for label in ('A.', 'B.', 'C.', 'D.'):
            self.assertIn(label, line)
        self.assertIn('\t', line, '選項之間要用 tab 分欄')
        self.assertEqual(sum(1 for c in options[0]._p if c.tag.endswith('}oMath')), 0,
                         f'選項行不該帶著範本的方程式：{line}')

    def test_head_block_is_centered(self):
        """卷頭整塊置中：卷名、副標題、姓名班別行（範本靠左也要拉回中間）。"""
        template = self.call('POST', '/api/paper-templates',
                             {'name': '副標題靠左的範本',
                              'file': data_url(make_docx_with_left_subtitle())})['template']
        out = self.call('POST', f"/api/paper-templates/{template['id']}/apply",
                        {'file': data_url(content_with_options())})
        doc = Document(io.BytesIO(base64.b64decode(out['docx'].split(',', 1)[1])))
        subtitle = next(p for p in doc.paragraphs if '禁止使用計算機' in p.text)
        self.assertEqual(subtitle.paragraph_format.alignment, WD_ALIGN_PARAGRAPH.CENTER)
        title = next(p for p in doc.paragraphs if '小測' in p.text)
        self.assertEqual(title.paragraph_format.alignment, WD_ALIGN_PARAGRAPH.CENTER)

    def test_math_can_be_exported_as_latex_code(self):
        """勾「數學符號用 LaTeX 代碼」：公式不再轉成 Word 方程式，直接留 $...$。"""
        template = self.upload()['template']
        out = self.call('POST', f"/api/paper-templates/{template['id']}/apply",
                        {'file': data_url(content_docx()), 'mathAsText': True})
        doc = Document(io.BytesIO(base64.b64decode(out['docx'].split(',', 1)[1])))
        equations = sum(1 for p in doc.paragraphs for c in p._p if c.tag.endswith('}oMath'))
        self.assertEqual(equations, 0, '勾了 LaTeX 代碼就不該有 Word 方程式')
        texts = [p.text for p in doc.paragraphs]
        self.assertTrue(any('$x^{2}-5x+6=0$' in t for t in texts), texts)
        self.assertTrue(any(r'$\frac{1}{2}+\frac{1}{3}$' in t for t in texts), texts)

    def test_numbering_only_paragraph_is_not_a_question(self):
        """「2.」單獨一行（整題是圖）不該變成空白題，否則存檔會被擋。"""
        doc = Document()
        doc.add_paragraph('一、單項選擇題')
        doc.add_paragraph('1.　題目一')
        doc.add_paragraph('2.')
        doc.add_paragraph('3.　題目三')
        buf = io.BytesIO()
        doc.save(buf)
        _, sections = paper_format.parse_content_bytes(buf.getvalue())
        self.assertEqual([q.stem for q in sections[0].questions], ['題目一', '題目三'])

    def test_rejects_non_docx(self):
        self.call('POST', '/api/paper-templates',
                  {'name': '壞檔', 'file': data_url(b'not a zip')}, expect=400)

    def test_student_cannot_use_templates(self):
        if not self.student:
            self.skipTest('沒有學生帳號')
        self.call('GET', '/api/paper-templates', user=self.student, expect=403)

    def test_delete_removes_template(self):
        template = self.upload()['template']
        self.call('DELETE', f"/api/paper-templates/{template['id']}")
        listed = self.call('GET', '/api/paper-templates')
        self.assertNotIn(template['id'], [t['id'] for t in listed['templates']])


    def test_to_tex_keeps_math_and_structure(self):
        out = self.call('POST', '/api/papers/to-tex', {'file': data_url(content_docx())})
        tex = out['tex']
        self.assertIn(r'\documentclass', tex)
        self.assertIn(r'\begin{enumerate}', tex)
        self.assertIn('$x^{2}-5x+6=0$', tex, '公式要原樣輸出，不能再轉義')
        self.assertIn(r'\frac{1}{2}+\frac{1}{3}', tex)
        self.assertIn('(本題目卷共兩頁，禁止使用計算機)', tex, '副標題不能掉')
        self.assertIn('禁止使用計算機', tex.split(r'\end{center}')[0],
                      '副標題要跟卷名一起置中')
        self.assertEqual(out['stats']['questions'], 2)
        self.assertTrue(out['filename'].endswith('.tex'))

    def test_to_tex_uses_custom_preamble(self):
        preamble = '\\documentclass{article}\n\\usepackage{amsmath}'
        out = self.call('POST', '/api/papers/to-tex',
                        {'file': data_url(content_docx()), 'preamble': preamble})
        self.assertTrue(out['tex'].startswith(preamble), out['tex'][:60])

    def test_to_tex_rejects_document_begin_in_preamble(self):
        self.call('POST', '/api/papers/to-tex',
                  {'file': data_url(content_docx()),
                   'preamble': '\\documentclass{article}\n\\begin{document}'}, expect=400)

    def test_student_cannot_export_tex(self):
        if not self.student:
            self.skipTest('沒有學生帳號')
        self.call('POST', '/api/papers/to-tex',
                  {'file': data_url(content_docx())}, user=self.student, expect=403)


PDF_LAYOUT = {
    'page': {'widthCm': 21, 'heightCm': 29.7,
             'marginsCm': {'top': 1.3, 'bottom': 1.3, 'left': 1.3, 'right': 1.3}},
    'headerLines': ['某某中學'],
    'title': {'sizePt': 17, 'bold': True},
    'subtitleLines': ['(禁止使用計算機)'],
    'infoLine': '高一＿＿班　姓名：＿＿＿＿',
    'section': {'sizePt': 12, 'bold': True, 'align': 'left'},
    'question': {'sizePt': 12, 'indentCm': 0},
    'options': {'perLine': 2, 'indentCm': 0.6, 'columnCm': 9.3},
    'marksPosition': 'leading',
    'answerLines': 0,
    'footerLines': [],
}

PDF_CONTENT = {
    'title': 'PDF 轉寫測試卷',
    'sections': [{
        'title': '一、計算題',
        'questions': [
            {'stem': '解方程 $x^{2}-5x+6=0$。', 'marks': 6, 'options': [], 'lines': []},
            {'stem': '計算 $\\frac{1}{2}+\\frac{1}{3}$。', 'marks': 4,
             'options': ['A. $\\frac{5}{6}$', 'B. $\\frac{1}{6}$'], 'lines': []},
        ],
    }],
}

FAKE_PDF = b'%PDF-1.4\n% fake pdf for tests\n'


class PdfImportTests(unittest.TestCase):
    """PDF 沒有結構，靠 AI 讀版面與內容；這裡用假模型驗證管線與權限。"""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.app = server.Application(self.tmp.name, server.ROOT)
        self.app.initialize()
        with self.app.db() as db:
            teacher = db.execute("SELECT * FROM users WHERE role='teacher' ORDER BY id LIMIT 1").fetchone()
        self.teacher = dict(teacher) if teacher else None
        if not self.teacher:
            self.skipTest('沒有教師帳號')

        def fake_call(instruction, payload, image=None, model=None, temperature=None, timeout=None):
            if '排版分析器' in instruction:
                return json.dumps(PDF_LAYOUT, ensure_ascii=False)
            if '轉寫器' in instruction:
                return json.dumps(PDF_CONTENT, ensure_ascii=False)
            return '{}'

        self.app.ai.call = fake_call
        # 同意機制只在「外部模型」模式下生效；測試要模擬外部，否則驗不到 403。
        self.app.ai.mode = 'external'
        self.patch = patch.object(server, 'pdf_page_images',
                                  return_value=['data:image/png;base64,AAAA'])
        self.patch.start()

    def tearDown(self):
        self.patch.stop()
        self.tmp.cleanup()

    def call(self, method, path, body=None, expect=200):
        try:
            payload, _ = self.app.dispatch(method, path, body or {}, self.teacher, None, '127.0.0.1')
        except server.APIError as exc:
            self.assertEqual(exc.status, expect, exc.message)
            return {'error': exc.message}
        self.assertEqual(expect, 200)
        return payload

    @staticmethod
    def pdf_url():
        return ('data:application/pdf;base64,'
                + base64.b64encode(FAKE_PDF).decode())

    def test_pdf_template_uses_ai_layout(self):
        out = self.call('POST', '/api/paper-templates',
                        {'name': 'PDF 範本', 'file': self.pdf_url(), 'consentExternal': True})
        spec = out['template']['spec']
        self.assertTrue(spec.get('aiLayout'))
        self.assertEqual(spec['source'], 'pdf')
        self.assertEqual(spec['marksPosition'], 'leading')
        self.assertEqual(spec['optionsPerLine'], 2)

    def test_pdf_template_needs_consent(self):
        self.call('POST', '/api/paper-templates',
                  {'name': 'PDF 範本', 'file': self.pdf_url()}, expect=403)

    def test_pdf_content_is_rendered(self):
        template = self.call('POST', '/api/paper-templates',
                             {'name': 'PDF 範本', 'file': self.pdf_url(),
                              'consentExternal': True})['template']
        out = self.call('POST', f"/api/paper-templates/{template['id']}/apply",
                        {'file': self.pdf_url(), 'consentExternal': True})
        self.assertEqual(out['stats']['questions'], 2)
        self.assertEqual(out['stats']['formulas'], 4)  # 第1題1條＋第2題題幹1條＋選項2條
        raw = base64.b64decode(out['docx'].split(',', 1)[1])
        doc = Document(io.BytesIO(raw))
        self.assertEqual(doc.paragraphs[0].text, 'PDF 轉寫測試卷')
        equations = sum(1 for p in doc.paragraphs for c in p._p if c.tag.endswith('}oMath'))
        self.assertEqual(equations, 4)

    def test_large_file_is_not_counted_as_text(self):
        """base64 的 PDF 很大，不該撞到「文字合計 80,000 字」的上限。"""
        big = b'%PDF-1.4\n%' + b'x' * 300_000
        out = self.call('POST', '/api/papers/to-tex',
                        {'file': 'data:application/pdf;base64,' + base64.b64encode(big).decode(),
                         'consentExternal': True})
        self.assertEqual(out['stats']['questions'], 2)

    def test_pdf_to_tex(self):
        out = self.call('POST', '/api/papers/to-tex',
                        {'file': self.pdf_url(), 'consentExternal': True})
        self.assertIn(r'\section*{一、計算題}', out['tex'])
        self.assertIn('$x^{2}-5x+6=0$', out['tex'])


class LatexExportTests(unittest.TestCase):
    """LaTeX 輸出：選項打橫，公式一律是能編譯的 LaTeX。"""

    def build(self, options, stem='求 $x$。', marks=None):
        section = paper_format.Section(
            title='一、選擇題', questions=[
                paper_format.Question(stem=stem, marks=marks, options=options)])
        return latex_export.to_tex('測試卷', [section], marks_position='trailing')

    def test_options_are_laid_out_across_one_line(self):
        tex = self.build(['A. $1$', 'B. $2$', 'C. $3$', 'D. $4$'])
        line = next(l for l in tex.splitlines() if 'A. ' in l and 'D. ' in l)
        for label in ('A. ', 'B. ', 'C. ', 'D. '):
            self.assertIn(label, line)
        self.assertIn(r'\makebox', line)
        self.assertNotIn(r'\begin{enumerate}[label=\Alph*', tex, '不要再直排')

    def test_long_options_fall_back_to_two_per_line(self):
        long = ['A. ' + '甲乙丙丁戊己庚辛壬癸' * 2, 'B. ' + '甲乙丙丁戊己庚辛壬癸' * 2]
        tex = self.build(long)
        line = next(l for l in tex.splitlines() if 'A. ' in l)
        self.assertNotIn('B. ', line, '太長就不要硬塞同一行')

    def test_degree_and_units_become_real_latex(self):
        tex = self.build(['A. $30\\degree$', 'B. 90$\\degree$'])
        self.assertNotIn(r'\degree', tex, '\\degree 沒定義，編不過')
        self.assertIn(r'$30^{\circ}$', tex)
        self.assertIn(r'$90^{\circ}$', tex, '數字要跟度數一起進公式')

    def test_fullwidth_punctuation_inside_math(self):
        tex = self.build(['A. $\\left（ \\rho ，\\theta \\right）$'])
        self.assertIn(r'$\left( \rho ,\theta \right)$', tex)

    def test_empty_root_index_and_functions(self):
        tex = self.build(['A. $\\sqrt[\\text{ }]{3}$', 'B. $y=\\rho sin\\theta$'])
        self.assertIn(r'$\sqrt{3}$', tex)
        self.assertIn(r'\sin\theta', tex)

    def test_text_wrapper_around_a_formula_is_removed(self):
        tex = self.build(['A. $\\text{4x+3y-3=0 }$', 'B. $-1\\text{或}5$'])
        self.assertIn('$4x+3y-3=0$', tex)
        self.assertIn(r'\text{或}', tex, '中文要留在 \\text 裡')

    def test_head_block_is_centered(self):
        section = paper_format.Section(title='一、選擇題', questions=[
            paper_format.Question(stem='求 $x$。', options=['A. $1$'])])
        tex = latex_export.to_tex('卷名', [section],
                                  subtitle_lines=['(禁止使用計算機)', '高一＿班　姓名：＿＿'])
        head = tex.split(r'\end{center}', 1)[0]
        self.assertIn('禁止使用計算機', head, '副標題要跟卷名一起置中')
        self.assertIn('姓名', head)
        self.assertIn(r'{\small ', head)

    def test_fill_in_blanks_are_solid_rules(self):
        tex = self.build(['A. 1'], stem='斜率 $k=$____________。')
        self.assertIn(r'\rule', tex)
        self.assertNotIn(r'\_\_\_', tex)


PAPER_SECTIONS = [{
    'title': '一、單項選擇題：（@3分，共30分）',
    'questions': [
        {'stem': '若 $l_{1}$ 的傾斜角為 $30\\degree$，則 $l_{1}$ 的斜率為(　　)', 'marks': 3,
         'options': ['A. $\\sqrt{3}$', 'B. $-\\sqrt{3}$',
                     'C. $\\frac{\\sqrt{3}}{3}$', 'D. $-\\frac{\\sqrt{3}}{3}$'], 'lines': []},
        {'stem': '直線的斜率 $k=$____________。', 'marks': 2, 'options': [], 'lines': []},
    ],
}]


class PaperLibraryTests(unittest.TestCase):
    """我的試卷：存起來、搜尋篩選、一鍵匯入題庫。"""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.app = server.Application(self.tmp.name, server.ROOT)
        self.app.initialize()
        with self.app.db() as db:
            teacher = db.execute("SELECT * FROM users WHERE role='teacher' ORDER BY id LIMIT 1").fetchone()
            student = db.execute("SELECT * FROM users WHERE role='student' ORDER BY id LIMIT 1").fetchone()
        self.teacher = dict(teacher) if teacher else None
        self.student = dict(student) if student else None
        if not self.teacher:
            self.skipTest('沒有教師帳號')

    def tearDown(self):
        self.tmp.cleanup()

    def call(self, method, path, body=None, user=None, expect=200, query=''):
        try:
            payload, _ = self.app.dispatch(method, path, body or {}, user or self.teacher,
                                           None, '127.0.0.1', query)
        except server.APIError as exc:
            self.assertEqual(exc.status, expect, exc.message)
            return {'error': exc.message}
        self.assertEqual(expect, 200)
        return payload

    def save(self, **override):
        body = {'title': '培華高二大測一', 'year': '2026', 'grade': '高二',
                'term': '第一段', 'sections': PAPER_SECTIONS}
        body.update(override)
        return self.call('POST', '/api/papers', body)['paper']

    def test_save_and_list(self):
        paper = self.save()
        self.assertEqual(paper['questionCount'], 2)
        self.assertEqual(paper['year'], '2026')
        listed = self.call('GET', '/api/papers')
        self.assertEqual([p['id'] for p in listed['papers']], [paper['id']])
        self.assertEqual(listed['facets'], {'years': ['2026'], 'grades': ['高二'], 'terms': ['第一段']})

    def test_search_and_filters(self):
        first = self.save()
        second = self.save(title='培華初三小測', year='2025', grade='初三', term='第二段',
                           sections=[{'title': '二、填空題', 'questions': [
                               {'stem': '計算 $1+1=$______。', 'marks': 2,
                                'options': [], 'lines': []}]}])
        by_year = self.call('GET', '/api/papers', query=urlencode({'year': '2026'}))
        self.assertEqual([p['id'] for p in by_year['papers']], [first['id']])
        by_grade = self.call('GET', '/api/papers', query=urlencode({'grade': '初三'}))
        self.assertEqual([p['id'] for p in by_grade['papers']], [second['id']])
        by_term = self.call('GET', '/api/papers', query=urlencode({'term': '第二段'}))
        self.assertEqual([p['id'] for p in by_term['papers']], [second['id']])
        # 搜尋也會比對題目文字（不只是卷名）
        by_text = self.call('GET', '/api/papers', query=urlencode({'q': '傾斜角'}))
        self.assertEqual([p['id'] for p in by_text['papers']], [first['id']])

    def test_detail_and_delete(self):
        paper = self.save()
        detail = self.call('GET', f"/api/papers/{paper['id']}")['paper']
        self.assertEqual(len(detail['sections'][0]['questions']), 2)
        self.call('DELETE', f"/api/papers/{paper['id']}")
        self.assertEqual(self.call('GET', '/api/papers')['papers'], [])

    def test_students_cannot_use_the_library(self):
        if not self.student:
            self.skipTest('沒有學生帳號')
        self.call('GET', '/api/papers', user=self.student, expect=403)
        self.call('POST', '/api/papers', {'title': 'x', 'sections': PAPER_SECTIONS},
                  user=self.student, expect=403)

    def test_one_click_import_to_bank(self):
        before = len(self.call('GET', '/api/state')['questions'])
        paper = self.save()
        out = self.call('POST', f"/api/papers/{paper['id']}/to-bank",
                        {'topic': '坐標幾何', 'grade': '中五', 'level': 2})
        self.assertEqual(out['count'], 2)
        choice, blank = out['questions']
        self.assertEqual(choice['qtype'], '選擇題')
        self.assertEqual(len(choice['options']), 4)
        self.assertEqual(choice['options'][0], '$\\sqrt{3}$', '選項的 A. 前綴要拆掉')
        self.assertEqual(choice['points'], 3)
        self.assertEqual(choice['answer'], '', '答案留空，等老師在題庫補')
        self.assertEqual(choice['grade'], '中五')
        self.assertEqual(choice['topic'], '坐標幾何')
        self.assertIn('待核驗', choice['tags'], '答案留空，要標成待核驗')
        self.assertIn('培華高二大測一', choice['source'])
        self.assertEqual(blank['qtype'], '填空題')
        self.assertEqual(blank['points'], 2)
        # 匯入後要記住，列表上看得到
        listed = self.call('GET', '/api/papers')['papers'][0]
        self.assertEqual(listed['importedCount'], 2)
        self.assertTrue(listed['importedAt'])
        # 題庫真的多了這兩題
        bank = self.call('GET', '/api/state')
        self.assertEqual(len(bank['questions']), before + 2, '匯入後題庫要真的多出這兩題')

    def test_save_skips_questions_without_text(self):
        paper = self.call('POST', '/api/papers', {
            'title': '有空題的卷',
            'sections': [{'title': '一、選擇題', 'questions': [
                {'stem': '有題目', 'marks': 2, 'options': [], 'lines': []},
                {'stem': '   ', 'marks': 2, 'options': [], 'lines': []}]}]})['paper']
        self.assertEqual(paper['questionCount'], 1)

    def test_import_rejects_unknown_topic(self):
        paper = self.save()
        self.call('POST', f"/api/papers/{paper['id']}/to-bank",
                  {'topic': '不存在的分類'}, expect=400)

    def test_saved_paper_exports_tex(self):
        paper = self.save()
        out = self.call('POST', f"/api/papers/{paper['id']}/to-tex", {})
        self.assertIn(r'\documentclass', out['tex'])
        self.assertIn('$30^{\\circ}$', out['tex'])
        self.assertIn(r'\makebox', out['tex'], '選項要打橫')
        self.assertTrue(out['filename'].endswith('.tex'))

    def test_apply_returns_content_for_saving(self):
        """排版完直接可以存：apply 要把題目內容帶回來，前端不用再上傳一次。"""
        template = self.call('POST', '/api/paper-templates',
                             {'name': '範本',
                              'file': data_url(make_docx(TEMPLATE))})['template']
        out = self.call('POST', f"/api/paper-templates/{template['id']}/apply",
                        {'file': data_url(content_docx())})
        content = out['content']
        self.assertEqual(content['title'], '中三數學 小測')
        self.assertEqual(len(content['sections']), 1)
        self.assertEqual(len(content['sections'][0]['questions']), 2)
        saved = self.call('POST', '/api/papers',
                          {'title': content['title'], 'year': '2026',
                           'grade': '初三', 'term': '第一段',
                           'sections': content['sections']})['paper']
        self.assertEqual(saved['questionCount'], 2)


if __name__ == '__main__':
    unittest.main()
