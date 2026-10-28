"""把新卷的內容套進範本的版式。

核心想法：**不重建排版，而是拿範本當容器**。
範本裡的頁面設定、頁眉頁腳、樣式、題號定位、答題橫線全部原封不動，
只把「題幹」與「配分」兩段文字換成新卷的內容——這樣才可能做到「完全一致」。

用法：
    python paper_format.py analyze 範本.docx
    python paper_format.py render 範本.docx 新卷.docx 輸出.docx
    python paper_format.py render-json 範本.docx 內容.json 輸出.docx
"""
from __future__ import annotations

import copy
import io
import json
import re
import sys
from dataclasses import dataclass, field, asdict
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Emu
from docx.text.paragraph import Paragraph

import omml

# ── 角色判定 ──────────────────────────────────────────────────────────
SECTION_RE = re.compile(r'^\s*(甲部|乙部|丙部|丁部|第[一二三四]部分|Part\s*[A-D]|[一二三四五六七八九十]+、)')
QUESTION_RE = re.compile(r'^\s*(\d{1,3})\s*[.、)．]')
OPTION_RE = re.compile(r'[(（]?([A-Da-d])\s*[.、)．]')
MARKS_RE = re.compile(r'[（(]\s*(\d+(?:\.\d+)?)\s*分\s*[)）]')
# 一行裡有多個選項（「A. … B. …」）時，要拆成獨立選項
OPTION_SPLIT_RE = re.compile(r'(?=\s*[(（]?[A-Da-d]\s*[.、)．])')


def split_options(text: str) -> list[str]:
    parts = [p.strip() for p in OPTION_SPLIT_RE.split(text) if p.strip()]
    return parts if len(parts) > 1 else [text.strip()]


def has_bottom_border(p) -> bool:
    ppr = p._p.pPr
    if ppr is None:
        return False
    pbdr = ppr.find(qn('w:pBdr'))
    if pbdr is None:
        return False
    bottom = pbdr.find(qn('w:bottom'))
    return bottom is not None and bottom.get(qn('w:val')) not in (None, 'nil', 'none')


def classify(text: str, p) -> str:
    t = text.strip()
    if not t:
        return 'answer_line' if has_bottom_border(p) else 'blank'
    if SECTION_RE.match(t):
        return 'section'
    if QUESTION_RE.match(t):
        return 'question'
    if OPTION_RE.match(t):
        return 'option'
    return 'other'


def run_style(p) -> dict:
    for r in p.runs:
        if r.text.strip():
            return {'font': r.font.name,
                    'sizePt': r.font.size.pt if r.font.size else None,
                    'bold': bool(r.font.bold)}
    return {}


def alignment_name(p) -> str:
    a = p.paragraph_format.alignment
    return {WD_ALIGN_PARAGRAPH.CENTER: 'center', WD_ALIGN_PARAGRAPH.RIGHT: 'right',
            WD_ALIGN_PARAGRAPH.JUSTIFY: 'justify', WD_ALIGN_PARAGRAPH.LEFT: 'left',
            None: 'left'}.get(a, str(a))


@dataclass
class Block:
    role: str
    text: str
    style: dict = field(default_factory=dict)
    align: str = 'left'
    lines_after: int = 0


def read_blocks(doc) -> list[Block]:
    """把段落分成角色。空的框線段落在題目後面才算「答題橫線」，否則算「分隔線」。"""
    blocks: list[Block] = []
    last_role = None
    for p in doc.paragraphs:
        role = classify(p.text, p)
        if role == 'answer_line' and last_role not in ('question', 'answer_line'):
            role = 'rule'          # 卷頭那條橫線，不是答題空間
        blocks.append(Block(role=role, text=p.text, style=run_style(p), align=alignment_name(p)))
        if role != 'blank':
            last_role = role
    for i, b in enumerate(blocks):
        if b.role != 'question':
            continue
        n = 0
        for nxt in blocks[i + 1:]:
            if nxt.role != 'answer_line':
                break
            n += 1
        b.lines_after = n
    return blocks


def marks_position(question_text: str) -> str:
    """配分寫在題號後面（leading）、推到右邊（trailing），還是根本沒有。"""
    if MARKS_RE.search(question_text[:16]):
        return 'leading'
    if '\t' in question_text and MARKS_RE.search(question_text.split('\t')[-1]):
        return 'trailing'
    return 'none'


def _analyze_doc(doc, source_name: str) -> dict:
    sec = doc.sections[0]
    blocks = read_blocks(doc)

    prototypes = {}
    for role in ('section', 'question', 'option', 'answer_line', 'rule', 'other'):
        for b in blocks:
            if b.role == role:
                prototypes[role] = asdict(b)
                break

    question_text = prototypes.get('question', {}).get('text', '')
    option_para = next((p for p in doc.paragraphs if classify(p.text, p) == 'option'), None)

    lines_by_index: list[int | None] = []
    for b in blocks:
        if b.role == 'section':
            lines_by_index.append(None)
        elif b.role == 'question' and lines_by_index and lines_by_index[-1] is None:
            lines_by_index[-1] = b.lines_after

    return {
        'source': source_name,
        'page': {
            'widthCm': round(sec.page_width.cm, 2), 'heightCm': round(sec.page_height.cm, 2),
            'marginsCm': {'top': round(sec.top_margin.cm, 2), 'bottom': round(sec.bottom_margin.cm, 2),
                          'left': round(sec.left_margin.cm, 2), 'right': round(sec.right_margin.cm, 2)},
            'headerCm': round(sec.header_distance.cm, 2), 'footerCm': round(sec.footer_distance.cm, 2),
        },
        'header': [p.text for p in sec.header.paragraphs if p.text.strip()],
        'footer': [p.text for p in sec.footer.paragraphs if p.text.strip()],
        'marksPosition': marks_position(question_text),
        'optionsPerLine': option_columns(option_para) if option_para is not None else 1,
        'prototypes': prototypes,
        'answerLinesBySection': lines_by_index,
        'counts': {r: sum(1 for b in blocks if b.role == r) for r in
                   ('section', 'question', 'option', 'answer_line', 'rule', 'other')},
        'tables': [{'rows': len(t.rows), 'cols': len(t.columns),
                    'firstCell': t.cell(0, 0).text.strip()[:40]} for t in doc.tables],
        'warnings': (['範本含表格：表格內容目前不會被帶進輸出'] if doc.tables else []),
    }


def analyze(path: str) -> dict:
    return _analyze_doc(Document(path), Path(path).name)


def analyze_bytes(data: bytes) -> dict:
    """伺服器上傳用：直接分析記憶體裡的 docx。"""
    return _analyze_doc(Document(io.BytesIO(data)), 'uploaded.docx')


# ── 內容模型 ──────────────────────────────────────────────────────────
@dataclass
class Question:
    stem: str
    marks: float | None = None
    options: list[str] = field(default_factory=list)
    lines: list[str] = field(default_factory=list)      # 小題或續行


@dataclass
class Section:
    title: str | None
    questions: list[Question] = field(default_factory=list)


def sections_from_json(path: str) -> tuple[str | None, list[Section], str | None]:
    data = json.loads(Path(path).read_text(encoding='utf-8'))
    sections = []
    for s in data.get('sections', []):
        sections.append(Section(title=s.get('title'), questions=[
            Question(stem=q['stem'], marks=q.get('marks'), options=q.get('options') or [],
                     lines=q.get('lines') or [])
            for q in s.get('questions', [])
        ]))
    return data.get('title'), sections, data.get('closing')


def _only_with_stem(sections: list[Section]) -> list[Section]:
    """題號後面沒有文字（整題是圖、或題號自己佔一行）的不要留成空白題。

    留著的話版面會多一個孤零零的「3.」，存到「我的試卷」也會被伺服器擋下來。
    """
    out = []
    for section in sections:
        questions = [q for q in section.questions if (q.stem or '').strip()]
        if questions:
            out.append(Section(title=section.title, questions=questions))
    return out


def parse_content(path: str) -> tuple[list[str], list[Section]]:
    doc = Document(path)
    preamble: list[str] = []
    sections: list[Section] = []
    current: Section | None = None
    for p in doc.paragraphs:
        text = omml.omml_latex_of_paragraph(p).strip() or p.text.strip()
        if not text:
            continue
        role = classify(text, p)
        if role == 'section':
            current = Section(title=text)
            sections.append(current)
        elif role == 'question':
            if current is None:
                current = Section(title=None)
                sections.append(current)
            m = MARKS_RE.search(text)
            marks = float(m.group(1)) if m else None
            stem = MARKS_RE.sub('', text)
            stem = QUESTION_RE.sub('', stem, count=1).strip()
            current.questions.append(Question(stem=stem, marks=marks))
        elif role == 'option' and current and current.questions:
            current.questions[-1].options.extend(split_options(text))
        elif current and current.questions:
            current.questions[-1].stem += text
        else:
            preamble.append(text)
    return preamble, _only_with_stem(sections)


# ── 套版 ──────────────────────────────────────────────────────────────
def _wipe(p):
    """清空段落內容，只留 pPr。

    以前的寫法是「只換文字 run」，可是範本的選項段落裡常常還夾著方程式
    （m:oMath）、tab、空白 run——那些不是文字 run，換不掉，結果每一行選項
    後面都拖著範本的範例公式。整段清掉重寫才不會有殘留。

    回傳第一個文字 run 的 rPr：新寫進去的文字要沿用它的字體與字級。
    """
    p_el = p._p
    rpr = None
    for r in p_el.findall(qn('w:r')):
        rpr_el = r.find(qn('w:rPr'))
        if rpr_el is None:
            continue
        text = ''.join(t.text or '' for t in r.findall(qn('w:t')))
        if text.strip():
            rpr = copy.deepcopy(rpr_el)
            break
    for child in list(p_el):
        if child.tag != qn('w:pPr'):
            p_el.remove(child)
    return rpr


def _run(p, text, rpr=None):
    run = p.add_run(text)
    if rpr is not None:
        run._r.insert(0, copy.deepcopy(rpr))
    return run


MATH_RE = re.compile(r'\$([^$]+)\$')


def split_math(text: str):
    """把「文字 $公式$ 文字」切成 [('text', ...), ('math', ...)]。"""
    parts, pos = [], 0
    for m in MATH_RE.finditer(text):
        if m.start() > pos:
            parts.append(('text', text[pos:m.start()]))
        parts.append(('math', m.group(1)))
        pos = m.end()
    if pos < len(text):
        parts.append(('text', text[pos:]))
    return parts or [('text', text)]


def _write(p, cells: list[str], rpr=None, math_as_text: bool = False):
    """把 cells 依序寫進（已清空的）段落：cell 之間補一個 tab，靠 pPr 的定位點分欄。

    math_as_text=True 時公式不轉成 Word 方程式，直接留下 $...$ 的 LaTeX 代碼。
    """
    for index, cell in enumerate(cells):
        if index:
            _run(p, '\t', rpr)
        for kind, value in split_math(cell or ''):
            if kind == 'math':
                if math_as_text:
                    _run(p, f'${value}$', rpr)
                else:
                    p._p.append(omml.latex_to_omml(value))
            elif value:
                _run(p, value, rpr)


def _fill(p, cells: list[str], math_as_text: bool = False):
    """清空後重寫，一次做完（範本殘留的公式／tab 不會漏出來）。"""
    _write(p, cells, _wipe(p), math_as_text)


def _tab_stops(p):
    """段落自己設的定位點；沒有就看它用的樣式。"""
    stops = list(p.paragraph_format.tab_stops)
    if stops:
        return stops
    try:
        return list(p.style.paragraph_format.tab_stops)
    except (AttributeError, KeyError):
        return []


def option_columns(p) -> int:
    """範本的選項一行排幾個：先看定位點，沒有定位點就數 A/B/C/D。"""
    stops = _tab_stops(p)
    if stops:
        return max(1, min(4, len(stops) + 1))
    return max(1, min(4, len(OPTION_RE.findall(p.text))))


def _text_width(doc, p):
    """段落可用的文字寬度（給「配分靠右」補定位點用）。"""
    sec = doc.sections[0]
    width = sec.page_width - sec.left_margin - sec.right_margin
    if p.paragraph_format.left_indent is not None:
        width -= p.paragraph_format.left_indent
    if p.paragraph_format.right_indent is not None:
        width -= p.paragraph_format.right_indent
    return width


# ── 選項打橫（一行 2～4 個）───────────────────────────────────────────
DEFAULT_OPTIONS_PER_LINE = 4
UNIT_CM = 0.23          # 12pt 半角字大約 0.21cm，留一點餘裕；全角算 2 個單位


def _units(text: str) -> float:
    """粗估一段文字的寬度（全角 2、半角 1），用來決定一行放得下幾個選項。"""
    plain = MATH_RE.sub(lambda m: 'x' * max(1, len(m.group(1)) // 2), text)
    return sum(2 if ord(ch) > 0x2E80 else 1 for ch in plain)


def fit_columns(options: list[str], per_line: int, width_cm: float) -> int:
    """一行放幾個：先用範本／預設的數量，選項太長就一路降。"""
    widest = max((_units(o) for o in options), default=1)
    for n in range(max(1, min(4, per_line)), 0, -1):
        if widest <= (width_cm / n) / UNIT_CM:
            return n
    return 1


def pack_options(options: list[str], columns: int) -> list[list[str]]:
    """把選項切成每行 columns 個。"""
    return [options[i:i + columns] for i in range(0, len(options), columns)]


def _column_stops(p, doc, columns: int):
    """範本沒有定位點時，自己按欄數補上（左邊界起算）。"""
    sec = doc.sections[0]
    full = sec.page_width - sec.left_margin - sec.right_margin
    indent = p.paragraph_format.left_indent or 0
    available = full - indent
    for i in range(1, columns):
        p.paragraph_format.tab_stops.add_tab_stop(
            Emu(int(indent + available * i / columns)), WD_TAB_ALIGNMENT.LEFT)


def _clone(doc, prototype_el, before):
    el = copy.deepcopy(prototype_el)
    before.addprevious(el)
    return Paragraph(el, doc)


def _marks_text(prototype_right: str, marks: float | None) -> str:
    if marks is None:
        return ''
    m = MARKS_RE.search(prototype_right or '')
    if not m:
        return f'（{marks:g} 分）'
    raw = m.group(0)
    space = ' ' if ' ' in raw else ''
    return f'{raw[0]}{marks:g}{space}分{raw[-1]}'


def _heading_text(template_title: str, new_title: str, questions) -> str:
    if not new_title or '（' in new_title or '(' in new_title:
        return new_title
    marks = [q.marks for q in questions if q.marks is not None]
    if not marks or not template_title:
        return new_title
    if template_title.strip()[:2] != new_title.strip()[:2] or '（' not in template_title:
        return new_title
    total = sum(marks)
    if len(set(marks)) == 1:
        return f'{new_title.strip()}（每題 {marks[0]:g} 分，共 {total:g} 分）'
    return f'{new_title.strip()}（共 {total:g} 分）'


def count_formulas(sections: list[Section]) -> int:
    total = 0
    for s in sections:
        for q in s.questions:
            for text in [q.stem, *q.lines, *q.options]:
                total += len(MATH_RE.findall(text))
    return total


def sections_to_dict(sections: list[Section]) -> list[dict]:
    """存檔用：Section/Question 轉成純 dict（我的試卷、匯入題庫都吃這個）。"""
    return [asdict(s) for s in sections]


def sections_from_dict(data) -> list[Section]:
    """存檔讀回來：純 dict → Section/Question（匯出 .tex 時用）。"""
    out = []
    for section in data or []:
        questions = [Question(stem=q.get('stem') or '', marks=q.get('marks'),
                              options=list(q.get('options') or []),
                              lines=list(q.get('lines') or []))
                     for q in section.get('questions') or []
                     if (q.get('stem') or '').strip()]
        if questions:
            out.append(Section(title=section.get('title'), questions=questions))
    return out


def _render_into(tpl, sections: list[Section], title: str | None = None,
                 closing: str | None = None,
                 options_per_line: int = DEFAULT_OPTIONS_PER_LINE,
                 math_as_text: bool = False) -> dict:
    """把內容套進一個已開啟的 Document（不存檔，讓呼叫端決定存哪裡）。"""
    blocks = read_blocks(tpl)
    paragraphs = tpl.paragraphs

    proto: dict[str, Paragraph] = {}
    for i, b in enumerate(blocks):
        if b.role not in proto:
            proto[b.role] = paragraphs[i]
    if 'question' not in proto:
        raise SystemExit('範本裡找不到題目段落（長得像「1. …」）')

    q_proto = proto['question']
    pos = marks_position(q_proto.text)
    marks_proto = q_proto.text.split('\t')[-1] if pos == 'trailing' else ''
    # 範本的「配分靠右」若只靠預設 tab，位置會飄；補一個貼右邊界的定位點。
    marks_tab = None
    if pos == 'trailing' and not _tab_stops(q_proto):
        marks_tab = _text_width(tpl, q_proto)

    # 卷頭：第一個分部之前的所有段落（標題、副標、資料行、橫線）
    first_section = next((i for i, b in enumerate(blocks) if b.role == 'section'), len(blocks))
    head_idx = [i for i in range(first_section) if blocks[i].role in ('other', 'rule', 'blank')]

    lines_by_index: list[int | None] = []
    for b in blocks:
        if b.role == 'section':
            lines_by_index.append(None)
        elif b.role == 'question' and lines_by_index and lines_by_index[-1] is None:
            lines_by_index[-1] = b.lines_after
    default_lines = next((n for n in lines_by_index if n), 0)

    body = tpl.element.body
    sectpr = body.find(qn('w:sectPr'))
    for child in list(body):
        if child is sectpr:
            continue
        body.remove(child)

    cloned_head = [_clone(tpl, paragraphs[i]._p, sectpr) for i in head_idx]
    # 卷頭整塊置中：卷名、副標題、姓名班別那幾行
    for p in cloned_head:
        if p.text.strip():
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if title:
        for p in cloned_head:
            if classify(p.text, p) == 'other' and p.runs:
                _fill(p, [title], math_as_text)
                break

    number = 1
    template_headings = [b.text for b in blocks if b.role == 'section']
    for index, s in enumerate(sections):
        if s.title and 'section' in proto:
            np = _clone(tpl, proto['section']._p, sectpr)
            tpl_title = template_headings[index] if index < len(template_headings) else ''
            _fill(np, [_heading_text(tpl_title, s.title, s.questions)], math_as_text)
        lines = lines_by_index[index] if index < len(lines_by_index) and lines_by_index[index] \
            else default_lines

        for q in s.questions:
            np = _clone(tpl, q_proto._p, sectpr)
            if pos == 'trailing':
                _fill(np, [f'{number}.　{q.stem}', _marks_text(marks_proto, q.marks)], math_as_text)
                if marks_tab is not None:
                    np.paragraph_format.tab_stops.add_tab_stop(marks_tab, WD_TAB_ALIGNMENT.RIGHT)
            elif q.marks is not None:
                _fill(np, [f'{number}.{_marks_text("（0 分）", q.marks)}　{q.stem}'], math_as_text)
            else:
                _fill(np, [f'{number}.　{q.stem}'], math_as_text)
            number += 1

            for extra in q.lines:
                ep = _clone(tpl, q_proto._p, sectpr)
                _fill(ep, [extra], math_as_text)

            if q.options and 'option' in proto:
                # 選項一律打橫排：範本有定位點就照範本的欄數，沒有就預設一行 4 個；
                # 選項太長會自動降成 3、2、1 個，不會疊字。
                per_line = option_columns(proto['option'])
                auto = per_line < 2
                if auto:
                    per_line = options_per_line
                width_cm = _text_width(tpl, proto['option']) / 360000
                columns = fit_columns(q.options, per_line, width_cm)
                for row in pack_options(q.options, columns):
                    op = _clone(tpl, proto['option']._p, sectpr)
                    _fill(op, row, math_as_text)
                    if auto and len(row) > 1:
                        _column_stops(op, tpl, columns)

            if 'answer_line' in proto and lines and not q.options:
                for _ in range(lines):
                    _clone(tpl, proto['answer_line']._p, sectpr)

    if closing and 'other' in proto:
        cp = _clone(tpl, proto['other']._p, sectpr)
        _fill(cp, [closing], math_as_text)

    if tpl.tables:
        print(f'⚠️ 範本有 {len(tpl.tables)} 個表格，內容未帶入輸出', file=sys.stderr)
    return {'title': title, 'sections': len(sections),
            'questions': sum(len(s.questions) for s in sections),
            'formulas': count_formulas(sections),
            'tables': len(tpl.tables)}


def render_sections(template_path: str, sections: list[Section], out_path: str,
                    title: str | None = None, closing: str | None = None,
                    options_per_line: int = DEFAULT_OPTIONS_PER_LINE,
                    math_as_text: bool = False) -> dict:
    tpl = Document(template_path)
    report = _render_into(tpl, sections, title, closing, options_per_line, math_as_text)
    tpl.save(out_path)
    report['out'] = str(Path(out_path).resolve())
    return report


def render(template_path: str, content_path: str, out_path: str) -> dict:
    preamble, sections = parse_content(content_path)
    return render_sections(template_path, sections, out_path,
                           title=preamble[0] if preamble else None)


def parse_content_bytes(data: bytes) -> tuple[list[str], list[Section]]:
    """跟 parse_content 一樣，但吃記憶體裡的 docx（伺服器上傳用）。"""
    doc = Document(io.BytesIO(data))
    preamble: list[str] = []
    sections: list[Section] = []
    current: Section | None = None
    for p in doc.paragraphs:
        text = omml.omml_latex_of_paragraph(p).strip() or p.text.strip()
        if not text:
            continue
        role = classify(text, p)
        if role == 'section':
            current = Section(title=text)
            sections.append(current)
        elif role == 'question':
            if current is None:
                current = Section(title=None)
                sections.append(current)
            m = MARKS_RE.search(text)
            marks = float(m.group(1)) if m else None
            stem = MARKS_RE.sub('', text)
            stem = QUESTION_RE.sub('', stem, count=1).strip()
            current.questions.append(Question(stem=stem, marks=marks))
        elif role == 'option' and current and current.questions:
            current.questions[-1].options.extend(split_options(text))
        elif current and current.questions:
            current.questions[-1].stem += text
        else:
            preamble.append(text)
    return preamble, _only_with_stem(sections)


def render_with_sections(template_bytes: bytes, sections: list[Section],
                         title: str | None = None,
                         options_per_line: int = DEFAULT_OPTIONS_PER_LINE,
                         math_as_text: bool = False) -> tuple[bytes, dict]:
    """已經有 Section 物件時用這個（例如 PDF 由 AI 轉寫出來的內容）。"""
    tpl = Document(io.BytesIO(template_bytes))
    report = _render_into(tpl, sections, title, None, options_per_line, math_as_text)
    buf = io.BytesIO()
    tpl.save(buf)
    return buf.getvalue(), report


def render_bytes(template_bytes: bytes, content_bytes: bytes,
                 options_per_line: int = DEFAULT_OPTIONS_PER_LINE,
                 math_as_text: bool = False) -> tuple[bytes, dict]:
    """伺服器用：兩份 docx 的 bytes 進，排好版的 docx bytes 出。"""
    tpl = Document(io.BytesIO(template_bytes))
    preamble, sections = parse_content_bytes(content_bytes)
    report = _render_into(tpl, sections, preamble[0] if preamble else None,
                          None, options_per_line, math_as_text)
    buf = io.BytesIO()
    tpl.save(buf)
    return buf.getvalue(), report


if __name__ == '__main__':
    cmd = sys.argv[1] if len(sys.argv) > 1 else 'analyze'
    if cmd == 'analyze':
        print(json.dumps(analyze(sys.argv[2]), ensure_ascii=False, indent=2))
    elif cmd == 'render':
        print(json.dumps(render(sys.argv[2], sys.argv[3], sys.argv[4]), ensure_ascii=False, indent=2))
    elif cmd == 'render-json':
        t, s, c = sections_from_json(sys.argv[3])
        print(json.dumps(render_sections(sys.argv[2], s, sys.argv[4], title=t, closing=c),
                         ensure_ascii=False, indent=2))
    else:
        raise SystemExit('用法：analyze <範本> | render <範本> <新卷.docx> <輸出> | '
                         'render-json <範本> <內容.json> <輸出>')
